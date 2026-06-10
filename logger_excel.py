from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADERS = [
    "timestamp",
    "experiment_id",
    "dataset",
    "model",
    "loss",
    "accuracy",
    "macro_precision",
    "macro_recall",
    "f1_macro",
    "f1_weighted",
    "gmean",
    "training_seconds",
    "inference_seconds",
    "samples",
    "notes",
]


def _create_workbook(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "results"

    ws.append(HEADERS)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.alignment = Alignment(horizontal="center")

    wb.save(path)


def append_result(
    path: str,
    row: Dict[str, Any],
) -> None:
    if not Path(path).exists():
        _create_workbook(path)

    wb = load_workbook(path)
    ws = wb["results"]

    values = []

    for header in HEADERS:
        if header == "timestamp":
            values.append(
                row.get(
                    header,
                    datetime.now().isoformat(timespec="seconds"),
                )
            )
        else:
            values.append(row.get(header, ""))

    ws.append(values)

    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        )

        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = min(max(max_len + 2, 12), 40)

    f1_col = HEADERS.index("f1_macro") + 1
    best_row = None
    best_value = -1.0

    for row_idx in range(2, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=f1_col).value

        try:
            numeric = float(value)
        except (TypeError, ValueError):
            continue

        if numeric > best_value:
            best_value = numeric
            best_row = row_idx

    for row_idx in range(2, ws.max_row + 1):
        for cell in ws[row_idx]:
            cell.fill = PatternFill(
                fill_type=None
            )

    if best_row is not None:
        for cell in ws[best_row]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="C6EFCE",
            )

    wb.save(path)
